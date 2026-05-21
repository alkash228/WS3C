from __future__ import annotations

import cgi
import datetime as dt
import errno
import html
import http.server
import json
import os
import re
import shutil
import socket
import socketserver
import base64
import subprocess
import tempfile
import time
import uuid
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse, parse_qs
from urllib.parse import urljoin
from urllib.request import Request, urlopen

import cv2
import numpy as np

PORT = int(os.environ.get("WEB_SAMV_PORT", "8090"))
MAX_PORT_TRIES = int(os.environ.get("WEB_SAMV_PORT_TRIES", "32"))
API_PORT_HINT = int(os.environ.get("WEB_SAMV_API_PORT", "8000"))
API_HOST_HINT = os.environ.get("WEB_SAMV_API_HOST", "localhost").strip() or "localhost"
API_BASE = os.environ.get("WEB_SAMV_API_BASE", "").strip().rstrip("/")
JOB_TIMEOUT_SEC = int(os.environ.get("WEB_SAMV_JOB_TIMEOUT_SEC", "3600"))
POLL_INTERVAL_SEC = float(os.environ.get("WEB_SAMV_POLL_INTERVAL_SEC", "1.2"))
AN_MIN_INTERSECTION_PX = int(os.environ.get("WEB_SAMV_ANALYZER_MIN_INTERSECTION_PX", "80"))
AN_MIN_INTERSECTION_RATIO = float(os.environ.get("WEB_SAMV_ANALYZER_MIN_INTERSECTION_RATIO", "0.03"))
AN_CROP_PAD_RATIO = float(os.environ.get("WEB_SAMV_ANALYZER_CROP_PAD_RATIO", "0.12"))

ROOT = Path(__file__).resolve().parent
FOLDERS_ROOT = (ROOT / "storage" / "folders").resolve()
INF_ROOT = (ROOT / "INF").resolve()
INF_BUILDINGS_FILE = INF_ROOT / "buildings.json"
INF_CONTRACTORS_FILE = INF_ROOT / "contractors.json"
INF_CONTRACTS_FILE = INF_ROOT / "contracts.json"
_RUNTIME: dict[str, object] = {"web_port": PORT, "api_port": API_PORT_HINT, "lan_ipv4": []}

VIDEO_EXTS = {".mp4", ".mov", ".avi", ".mkv", ".webm", ".m4v"}
_INF_DEFAULTS = {
    "buildings": ["Корпус 1", "Корпус 2"],
    "contractors": ["ПАО Газпром", "ООО Рога и Копыта", "ООО БНАЛ"],
    "contracts": [
        "Договор подряда №1 от 01.01.2026",
        "Договор подряда №2 от 15.02.2026",
    ],
}


def _safe_name(raw: str) -> str:
    value = (raw or "").strip()
    value = re.sub(r"\s+", "_", value)
    value = re.sub(r"[^A-Za-z0-9_-]", "", value)
    return value.strip("._-")


def _ensure_dirs() -> None:
    FOLDERS_ROOT.mkdir(parents=True, exist_ok=True)
    INF_ROOT.mkdir(parents=True, exist_ok=True)
    if not INF_BUILDINGS_FILE.is_file():
        INF_BUILDINGS_FILE.write_text(
            json.dumps(_INF_DEFAULTS["buildings"], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    if not INF_CONTRACTORS_FILE.is_file():
        INF_CONTRACTORS_FILE.write_text(
            json.dumps(_INF_DEFAULTS["contractors"], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
    if not INF_CONTRACTS_FILE.is_file():
        INF_CONTRACTS_FILE.write_text(
            json.dumps(_INF_DEFAULTS["contracts"], ensure_ascii=False, indent=2),
            encoding="utf-8",
        )


def _norm_inf_item(raw: str) -> str:
    s = str(raw or "").strip()
    s = re.sub(r"\s+", " ", s)
    return s[:120].strip()


def _norm_inf_kind(kind: str) -> str:
    key = str(kind or "").strip().lower()
    aliases = {
        "building": "buildings",
        "contractor": "contractors",
        "contract": "contracts",
        "subcontract": "contracts",
        "subcontracts": "contracts",
    }
    return aliases.get(key, key)


def _kind_to_inf_path(kind: str) -> Path:
    key = _norm_inf_kind(kind)
    mapping = {
        "buildings": INF_BUILDINGS_FILE,
        "contractors": INF_CONTRACTORS_FILE,
        "contracts": INF_CONTRACTS_FILE,
    }
    path = mapping.get(key)
    if path is None:
        raise ValueError("kind must be 'buildings', 'contractors' or 'contracts'")
    return path


def _defaults_for_inf_path(path: Path) -> list[str]:
    if path == INF_BUILDINGS_FILE:
        return _INF_DEFAULTS["buildings"]
    if path == INF_CONTRACTORS_FILE:
        return _INF_DEFAULTS["contractors"]
    return _INF_DEFAULTS["contracts"]


def _read_inf_list(path: Path, defaults: list[str]) -> list[str]:
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        raw = []
    out: list[str] = []
    if isinstance(raw, list):
        for x in raw:
            s = _norm_inf_item(str(x or ""))
            if s and s not in out:
                out.append(s)
    if not out:
        out = [_norm_inf_item(x) for x in defaults if _norm_inf_item(x)]
    return out


def _write_inf_list(path: Path, values: list[str]) -> None:
    uniq: list[str] = []
    for x in values:
        s = _norm_inf_item(x)
        if s and s not in uniq:
            uniq.append(s)
    path.write_text(json.dumps(uniq, ensure_ascii=False, indent=2), encoding="utf-8")


def _inf_options() -> dict[str, list[str]]:
    _ensure_dirs()
    return {
        "inf_version": 2,
        "buildings": _read_inf_list(INF_BUILDINGS_FILE, _INF_DEFAULTS["buildings"]),
        "contractors": _read_inf_list(INF_CONTRACTORS_FILE, _INF_DEFAULTS["contractors"]),
        "contracts": _read_inf_list(INF_CONTRACTS_FILE, _INF_DEFAULTS["contracts"]),
    }


_REPORT_FORMATS = {"json", "txt", "html", "docx"}


def _norm_report_formats(raw: object) -> set[str]:
    if not isinstance(raw, list):
        return set(_REPORT_FORMATS)
    out: set[str] = set()
    aliases = {"word": "docx", "doc": "docx", "docx": "docx"}
    for x in raw:
        key = str(x or "").strip().lower()
        key = aliases.get(key, key)
        if key in _REPORT_FORMATS:
            out.add(key)
    return out or set(_REPORT_FORMATS)


def _analyzer_mid_warning(warnings: list[dict]) -> dict | None:
    if not warnings:
        return None
    idx = len(warnings) // 2
    x = warnings[idx]
    return x if isinstance(x, dict) else None


def _norm_report_resolved(raw: object) -> str:
    if isinstance(raw, bool):
        return "Устранено" if raw else "Не устранено"
    s = str(raw or "").strip()
    if not s:
        return ""
    low = s.lower()
    if low in {"1", "true", "yes", "да", "устранено", "resolved"}:
        return "Устранено"
    if low in {"0", "false", "no", "нет", "не устранено", "open"}:
        return "Не устранено"
    return s[:120]


def _norm_report_items(items_raw: object) -> list[dict[str, object]]:
    out: list[dict[str, object]] = []
    if not isinstance(items_raw, list):
        return out
    for x in items_raw:
        if not isinstance(x, dict):
            continue
        hid = str(x.get("human_id", "") or "").strip()
        building = _norm_inf_item(str(x.get("building", "") or ""))
        contractor = _norm_inf_item(str(x.get("contractor", "") or ""))
        contract = _norm_inf_item(str(x.get("contract", "") or ""))
        frame = int(x.get("frame", -1) or -1)
        image_url = str(x.get("image_url", "") or "").strip()
        violations = str(x.get("violations", "") or "").strip()
        deadline = str(x.get("deadline", "") or "").strip()[:32]
        resolved = _norm_report_resolved(x.get("resolved", ""))
        reasons_raw = x.get("reasons")
        reasons: list[str] = []
        if isinstance(reasons_raw, list):
            reasons = [str(r or "").strip() for r in reasons_raw if str(r or "").strip()]
        if not violations and reasons:
            violations = "\n".join(reasons)
        if not hid:
            continue
        if not building or not contractor or not contract:
            continue
        out.append(
            {
                "human_id": hid,
                "building": building,
                "contractor": contractor,
                "contract": contract,
                "frame": frame,
                "image_url": image_url,
                "violations": violations,
                "deadline": deadline,
                "resolved": resolved,
                "reasons": reasons,
            }
        )
    return out


def _report_image_path(image_url: str) -> Path | None:
    s = str(image_url or "").strip().split("?", 1)[0]
    prefix = "/storage/folders/"
    if not s.startswith(prefix):
        return None
    rel = s[len(prefix) :]
    p = (FOLDERS_ROOT / rel).resolve()
    try:
        p.relative_to(FOLDERS_ROOT)
    except ValueError:
        return None
    return p if p.is_file() else None


def _report_status_badge(resolved: str) -> str:
    s = str(resolved or "").strip()
    if not s:
        return '<span class="badge badge-muted">—</span>'
    low = s.lower()
    if low in {"устранено", "да", "yes", "resolved"}:
        cls = "badge-ok"
    elif low in {"не устранено", "нет", "no", "open"}:
        cls = "badge-warn"
    else:
        cls = "badge-info"
    return f'<span class="badge {cls}">{html.escape(s)}</span>'


def _report_html(report: dict) -> str:
    items = report.get("items")
    if not isinstance(items, list):
        items = []
    rows: list[str] = []
    for n, x in enumerate(items, start=1):
        if not isinstance(x, dict):
            continue
        hid = html.escape(str(x.get("human_id", "")))
        building = html.escape(str(x.get("building", "")))
        contractor = html.escape(str(x.get("contractor", "")))
        contract = html.escape(str(x.get("contract", "")))
        row_img_url = html.escape(str(x.get("image_url", "") or ""))
        violations = html.escape(str(x.get("violations", "") or "")).replace("\n", "<br/>")
        deadline = html.escape(str(x.get("deadline", "") or "")) or "—"
        resolved = _report_status_badge(str(x.get("resolved", "") or ""))
        if row_img_url:
            img_cell = (
                f'<div class="photo-wrap"><img class="photo" src="{row_img_url}" '
                f'alt="фото нарушения" /></div>'
            )
        else:
            img_cell = '<span class="muted">—</span>'
        row_cls = "row-even" if n % 2 == 0 else "row-odd"
        rows.append(
            f'<tr class="{row_cls}">'
            f'<td class="col-num">{n}</td>'
            f'<td class="col-id">{hid}</td>'
            f'<td class="col-photo">{img_cell}</td>'
            f"<td>{building}</td>"
            f"<td>{contractor}</td>"
            f"<td>{contract}</td>"
            f'<td class="col-viol">{violations or "—"}</td>'
            f'<td class="col-date">{deadline}</td>'
            f'<td class="col-status">{resolved}</td>'
            "</tr>"
        )
    table_html = "\n".join(rows) if rows else '<tr><td colspan="9" class="empty">Нет записей</td></tr>'
    gen_at = html.escape(str(report.get("generated_at", "")))
    folder = html.escape(str(report.get("folder", "")))
    rep_id = html.escape(str(report.get("id", "")))
    frames = html.escape(str(report.get("frames_checked", "")))
    warns = html.escape(str(report.get("warnings_count", "")))
    row_count = len(rows)
    return f"""<!doctype html>
<html lang="ru">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width,initial-scale=1" />
  <title>Отчёт {rep_id}</title>
  <style>
    * {{ box-sizing: border-box; }}
    body {{
      font: 14px/1.5 "Segoe UI", Arial, sans-serif;
      margin: 0;
      padding: 28px 20px 40px;
      background: linear-gradient(160deg, #e8eef8 0%, #f5f7fb 45%, #eef2f9 100%);
      color: #1a2332;
    }}
    .page {{ max-width: 1280px; margin: 0 auto; }}
    .header {{
      background: linear-gradient(135deg, #1e3a6e 0%, #2d5aa8 55%, #3d6fbf 100%);
      color: #fff;
      border-radius: 14px;
      padding: 22px 26px;
      margin-bottom: 20px;
      box-shadow: 0 8px 24px rgba(30, 58, 110, 0.22);
    }}
    .header h1 {{ margin: 0 0 6px; font-size: 22px; font-weight: 600; }}
    .header .sub {{ opacity: 0.9; font-size: 13px; }}
    .stats {{ display: flex; flex-wrap: wrap; gap: 12px; margin-top: 16px; }}
    .stat {{
      background: rgba(255,255,255,0.14);
      border: 1px solid rgba(255,255,255,0.22);
      border-radius: 10px;
      padding: 10px 14px;
      min-width: 140px;
    }}
    .stat .lbl {{ font-size: 11px; text-transform: uppercase; letter-spacing: 0.06em; opacity: 0.85; }}
    .stat .val {{ font-size: 15px; font-weight: 600; margin-top: 4px; }}
    .table-card {{
      background: #fff;
      border-radius: 14px;
      border: 1px solid #d4deef;
      box-shadow: 0 4px 18px rgba(26, 35, 50, 0.07);
      overflow: hidden;
    }}
    .table-title {{
      padding: 14px 20px;
      border-bottom: 1px solid #e4ebf6;
      font-size: 16px;
      font-weight: 600;
      color: #1e3a6e;
    }}
    .table-scroll {{ overflow-x: auto; }}
    table.report {{ width: 100%; border-collapse: collapse; font-size: 13px; }}
    table.report thead th {{
      background: #eef3fb;
      color: #1e3a6e;
      font-weight: 600;
      text-align: left;
      padding: 11px 10px;
      border-bottom: 2px solid #c5d4eb;
      white-space: nowrap;
    }}
    table.report tbody td {{
      padding: 10px;
      border-bottom: 1px solid #e8edf5;
      vertical-align: top;
    }}
    table.report tbody tr.row-even {{ background: #fafbfd; }}
    table.report tbody tr:hover {{ background: #f0f5fd; }}
    .col-num {{ width: 36px; text-align: center; color: #6b7a90; font-weight: 600; }}
    .col-id {{ width: 56px; font-weight: 600; color: #2d5aa8; }}
    .col-photo {{ width: 150px; text-align: center; }}
    .col-viol {{ min-width: 200px; max-width: 320px; }}
    .col-date {{ width: 110px; white-space: nowrap; }}
    .col-status {{ width: 130px; text-align: center; }}
    .photo-wrap {{
      display: inline-block;
      padding: 4px;
      background: #f4f7fb;
      border: 1px solid #d4deef;
      border-radius: 8px;
    }}
    .photo {{
      display: block;
      max-width: 128px;
      max-height: 96px;
      border-radius: 4px;
      object-fit: contain;
    }}
    .badge {{
      display: inline-block;
      padding: 4px 10px;
      border-radius: 999px;
      font-size: 12px;
      font-weight: 600;
    }}
    .badge-ok {{ background: #e6f6ec; color: #1a7a3e; border: 1px solid #b8e6c8; }}
    .badge-warn {{ background: #fff4e5; color: #9a5b00; border: 1px solid #ffd89a; }}
    .badge-info {{ background: #e8f0ff; color: #1e4a8c; border: 1px solid #c5d8f5; }}
    .badge-muted {{ background: #f0f2f5; color: #6b7a90; border: 1px solid #dde2ea; }}
    .muted {{ color: #8a96a8; }}
    td.empty {{ text-align: center; padding: 24px; color: #8a96a8; }}
    @media print {{
      body {{ background: #fff; padding: 0; }}
      .header {{ box-shadow: none; }}
      .table-card {{ box-shadow: none; }}
    }}
  </style>
</head>
<body>
  <div class="page">
    <header class="header">
      <h1>Отчёт о выявленных нарушениях</h1>
      <p class="sub">WEB_samv · {rep_id}</p>
      <div class="stats">
        <div class="stat"><div class="lbl">Сформирован</div><div class="val">{gen_at}</div></div>
        <div class="stat"><div class="lbl">Папка</div><div class="val">{folder}</div></div>
        <div class="stat"><div class="lbl">Проверено кадров</div><div class="val">{frames}</div></div>
        <div class="stat"><div class="lbl">WARNING</div><div class="val">{warns}</div></div>
        <div class="stat"><div class="lbl">Записей</div><div class="val">{row_count}</div></div>
      </div>
    </header>
    <section class="table-card">
      <div class="table-title">Реестр нарушений</div>
      <div class="table-scroll">
        <table class="report">
          <thead>
            <tr>
              <th>№</th>
              <th>ID</th>
              <th>Фото</th>
              <th>Здание</th>
              <th>Подрядчик</th>
              <th>Договор подряда</th>
              <th>Перечень нарушений</th>
              <th>Срок устранения</th>
              <th>Отметка об устранении</th>
            </tr>
          </thead>
          <tbody>
            {table_html}
          </tbody>
        </table>
      </div>
    </section>
  </div>
</body>
</html>
"""



def _report_docx_embed_image(cell, img_path: Path) -> None:
    try:
        from docx.shared import Cm
    except ImportError as exc:
        raise RuntimeError("python-docx не установлен (pip install python-docx)") from exc
    cell.text = ""
    p = cell.paragraphs[0]
    try:
        p.add_run().add_picture(str(img_path), width=Cm(2.8))
    except Exception:
        try:
            from PIL import Image

            with Image.open(img_path) as im:
                im = im.convert("RGB")
                tmp = img_path.parent / f"_docx_{img_path.stem}.jpg"
                im.save(tmp, format="JPEG", quality=85)
            p.add_run().add_picture(str(tmp), width=Cm(2.8))
        except Exception:
            p.add_run("—")


def _report_docx(report: dict, docx_path: Path) -> None:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError as exc:
        raise RuntimeError("python-docx не установлен (pip install python-docx)") from exc

    doc = Document()
    doc.add_heading("Отчёт о выявленных нарушениях", level=1)
    doc.add_paragraph(f"WEB_samv · {report.get('id', '')}")
    meta = [
        ("Сформирован", str(report.get("generated_at", ""))),
        ("Папка", str(report.get("folder", ""))),
        ("WARNING (всего)", str(report.get("warnings_count", ""))),
        ("Записей в отчёте", str(len(report.get("items") or []))),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)

    items = report.get("items")
    if not isinstance(items, list):
        items = []
    doc.add_paragraph()
    h = doc.add_paragraph("Реестр нарушений")
    if h.runs:
        h.runs[0].bold = True
        h.runs[0].font.size = Pt(13)

    headers = [
        "№",
        "ID",
        "Фото",
        "Здание",
        "Подрядчик",
        "Договор подряда",
        "Перечень нарушений",
        "Срок устранения",
        "Отметка",
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        for p in hdr[i].paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.size = Pt(9)

    for n, x in enumerate(items, start=1):
        if not isinstance(x, dict):
            continue
        row = table.add_row().cells
        row[0].text = str(n)
        row[1].text = str(x.get("human_id", ""))
        img_path = _report_image_path(str(x.get("image_url", "") or ""))
        if img_path:
            _report_docx_embed_image(row[2], img_path)
        else:
            row[2].text = "—"
        row[3].text = str(x.get("building", ""))
        row[4].text = str(x.get("contractor", ""))
        row[5].text = str(x.get("contract", ""))
        row[6].text = str(x.get("violations", "") or "—")
        row[7].text = str(x.get("deadline", "") or "—")
        row[8].text = str(x.get("resolved", "") or "—")

    doc.save(str(docx_path))


def _generate_report(
    folder: str,
    frame: int,
    image_url: str,
    items: list[dict[str, object]],
    formats: set[str] | None = None,
) -> dict[str, object]:
    folder_path, _, _ = _load_folder_payload(folder)
    analysis_dir = folder_path / "analysis"
    analysis_dir.mkdir(parents=True, exist_ok=True)
    result_path = analysis_dir / "analyzer_result.json"
    if not result_path.is_file():
        raise FileNotFoundError("Run analysis first")
    payload = json.loads(result_path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("Invalid analyzer_result.json")
    warnings = payload.get("warnings")
    if not isinstance(warnings, list):
        warnings = []
    mid = _analyzer_mid_warning(warnings)
    if frame < 0 and isinstance(mid, dict):
        frame = int(mid.get("frame", -1) or -1)
    if not image_url and isinstance(mid, dict):
        image_url = str(mid.get("image_url", "") or "")
    report_id = dt.datetime.now().strftime("report_%Y%m%d_%H%M%S")
    report = {
        "schema": "samv_report_v2",
        "id": report_id,
        "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        "folder": folder,
        "frames_checked": int(payload.get("frames_checked", 0) or 0),
        "warnings_count": int(payload.get("warnings_count", 0) or 0),
        "items": items,
    }
    json_name = f"{report_id}.json"
    txt_name = f"{report_id}.txt"
    html_name = f"{report_id}.html"
    docx_name = f"{report_id}.docx"
    json_path = analysis_dir / json_name
    txt_path = analysis_dir / txt_name
    html_path = analysis_dir / html_name
    docx_path = analysis_dir / docx_name
    fmts = formats if formats else set(_REPORT_FORMATS)
    txt_lines = [
        "WEB_samv report",
        f"id: {report_id}",
        f"generated_at: {report['generated_at']}",
        f"folder: {folder}",
        f"frames_checked: {report['frames_checked']}",
        f"warnings_count: {report['warnings_count']}",
    ]
    for x in items:
        txt_lines.extend(
            [
                f"--- human_id:{x.get('human_id')} ---",
                f"building: {x.get('building')}",
                f"contractor: {x.get('contractor')}",
                f"contract: {x.get('contract')}",
                f"deadline: {x.get('deadline', '')}",
                f"resolved: {x.get('resolved', '')}",
                f"image: {x.get('image_url', '')}",
                "violations:",
                str(x.get("violations", "") or ""),
            ]
        )
    out: dict[str, object] = {"ok": True, "report": report, "formats": sorted(fmts)}
    if "json" in fmts:
        json_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        out["report_json_url"] = f"/storage/folders/{folder}/analysis/{json_name}"
    if "txt" in fmts:
        txt_path.write_text("\n".join(txt_lines) + "\n", encoding="utf-8")
        out["report_txt_url"] = f"/storage/folders/{folder}/analysis/{txt_name}"
    if "html" in fmts:
        html_path.write_text(_report_html(report), encoding="utf-8")
        out["report_html_url"] = f"/storage/folders/{folder}/analysis/{html_name}"
    if "docx" in fmts:
        _report_docx(report, docx_path)
        out["report_docx_url"] = f"/storage/folders/{folder}/analysis/{docx_name}"
    return out


def _json_response(handler: http.server.SimpleHTTPRequestHandler, payload: dict, code: int = 200) -> None:
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    handler.send_response(code)
    handler.send_header("Content-Type", "application/json; charset=utf-8")
    handler.send_header("Cache-Control", "no-store")
    handler.send_header("Content-Length", str(len(body)))
    handler.end_headers()
    handler.wfile.write(body)


def _folder_paths(folder_name: str) -> tuple[Path, Path, Path]:
    folder = FOLDERS_ROOT / folder_name
    return folder, folder / "data.json", folder / "video"


def _find_video_any(folder: Path) -> Path | None:
    for p in folder.iterdir():
        if p.is_file() and p.suffix.lower() in VIDEO_EXTS and p.stem == "video":
            return p
    return None


def _load_folder_payload(folder_name: str) -> tuple[Path, dict, Path | None]:
    name = _safe_name(folder_name)
    if not name:
        raise ValueError("Invalid folder")
    folder = FOLDERS_ROOT / name
    if not folder.is_dir():
        raise FileNotFoundError("Folder not found")
    json_path = folder / "data.json"
    if not json_path.is_file():
        raise FileNotFoundError("data.json not found")
    payload = json.loads(json_path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("Invalid data.json payload")
    return folder, payload, _find_video_any(folder)


def _mask_from_rle_row_major(enc: dict, h: int, w: int) -> np.ndarray | None:
    if not isinstance(enc, dict):
        return None
    if str(enc.get("format", "")) != "rle_row_major":
        return None
    if str(enc.get("order", "C")) != "C":
        return None
    eh = int(enc.get("height", 0) or 0)
    ew = int(enc.get("width", 0) or 0)
    if eh != h or ew != w or h <= 0 or w <= 0:
        return None
    counts = enc.get("counts")
    if not isinstance(counts, list):
        return None
    flat = np.zeros((h * w,), dtype=np.uint8)
    pos = 0
    val = 0
    for x in counts:
        n = int(x or 0)
        if n > 0:
            end = min(flat.size, pos + n)
            if val == 1:
                flat[pos:end] = 1
            pos = end
        val = 1 - val
        if pos >= flat.size:
            break
    return flat.reshape((h, w)).astype(bool)


def _instances_for_frame(payload: dict, frame_idx: int) -> list[dict]:
    frames = payload.get("frames")
    if not isinstance(frames, list):
        return []
    for fr in frames:
        if not isinstance(fr, dict):
            continue
        if int(fr.get("frame", -1)) != int(frame_idx):
            continue
        inst = fr.get("instances")
        if isinstance(inst, list):
            return [x for x in inst if isinstance(x, dict)]
        return []
    return []


def _all_prompts(payload: dict) -> list[dict[str, object]]:
    seen: dict[str, int] = {}
    declared: list[str] = []
    raw_prompt = str(payload.get("prompt", "") or "").strip()
    if raw_prompt:
        parts = re.split(r"[,\n;]+|\s\.\s|\.", raw_prompt)
        for p in parts:
            s = str(p or "").strip()
            if s and s not in declared:
                declared.append(s)
    frames = payload.get("frames")
    if not isinstance(frames, list):
        return [{"label": x, "prompt_id": i + 1} for i, x in enumerate(declared)]
    for fr in frames:
        if not isinstance(fr, dict):
            continue
        inst = fr.get("instances")
        if not isinstance(inst, list):
            continue
        for x in inst:
            if not isinstance(x, dict):
                continue
            lbl = str(x.get("prompt_label", "") or "").strip()
            if not lbl:
                continue
            pid = x.get("prompt_id")
            if lbl not in seen:
                try:
                    seen[lbl] = int(pid)
                except Exception:
                    seen[lbl] = 0
    for i, lbl in enumerate(declared, start=1):
        if lbl not in seen:
            seen[lbl] = i
    out = [{"label": k, "prompt_id": int(v)} for k, v in seen.items()]
    out.sort(key=lambda r: str(r.get("label", "")))
    return out


def _label_slug(label: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(label or "").strip().lower()).strip("_")


def _label_segments(raw: str) -> list[str]:
    s = str(raw or "").strip()
    if not s:
        return []
    parts = [str(x or "").strip() for x in re.split(r"[,\n;]+|\s\.\s|\.", s) if str(x or "").strip()]
    if not parts:
        return [s]
    # Keep order but remove duplicates.
    seen: set[str] = set()
    out: list[str] = []
    for p in parts:
        k = p.casefold()
        if k in seen:
            continue
        seen.add(k)
        out.append(p)
    return out


def _label_matches(instance_label: str, wanted_label: str) -> bool:
    il = str(instance_label or "").strip()
    wl = str(wanted_label or "").strip()
    if not il or not wl:
        return False
    if il == wl:
        return True
    # Compatibility: some JSONs store one combined label like
    # "human, person head, helmet". Treat it as multiple segments.
    wanted_cf = wl.casefold()
    for seg in _label_segments(il):
        if seg.casefold() == wanted_cf:
            return True
    return False


def _instance_group_by_label(instances: list[dict], label: str, h: int, w: int) -> list[tuple[int, np.ndarray]]:
    out: list[tuple[int, np.ndarray]] = []
    slug = _label_slug(label)
    key = f"{slug}_id" if slug else ""
    for x in instances:
        lbl = str(x.get("prompt_label", "") or "").strip()
        if not _label_matches(lbl, label):
            continue
        m = _mask_from_rle_row_major(x.get("mask", {}), h, w)
        if m is None or not m.any():
            continue
        obj_id = None
        # В приоритете реальный инстансный object_id из трекера.
        # Это гарантирует раздельный анализ каждого объекта даже если *_id в JSON были агрегированными.
        if "object_id" in x:
            try:
                obj_id = int(x.get("object_id"))
            except Exception:
                obj_id = None
        if obj_id is None and key and key in x:
            try:
                obj_id = int(x.get(key))
            except Exception:
                obj_id = None
        if obj_id is None:
            try:
                obj_id = int(x.get("prompt_id"))
            except Exception:
                obj_id = None
        if obj_id is None:
            try:
                obj_id = int(x.get("object_id"))
            except Exception:
                obj_id = 0
        out.append((int(obj_id), m))
    return out


def _bbox_from_mask(mask: np.ndarray) -> tuple[int, int, int, int] | None:
    ys, xs = np.where(mask)
    if ys.size == 0 or xs.size == 0:
        return None
    x0, x1 = int(xs.min()), int(xs.max())
    y0, y1 = int(ys.min()), int(ys.max())
    return x0, y0, x1 - x0 + 1, y1 - y0 + 1


def _mask_overlap_ok(a: np.ndarray, b: np.ndarray) -> bool:
    inter = np.logical_and(a, b)
    inter_px = int(inter.sum())
    if inter_px <= 0:
        return False
    if inter_px < int(AN_MIN_INTERSECTION_PX):
        return False
    b_area = int(b.sum())
    if b_area <= 0:
        return False
    ratio = float(inter_px) / float(b_area)
    return ratio >= float(AN_MIN_INTERSECTION_RATIO)


def _mask_intersects(a: np.ndarray, b: np.ndarray) -> bool:
    return bool(np.logical_and(a, b).any())


def _reason_for_overlay(reason: str) -> str:
    s = str(reason or "").strip()
    if "->" in s and "нет " in s and " на кадре" in s:
        left, right = s.split("->", 1)
        dep = right.replace("нет ", "").replace(" на кадре", "").strip()
        return f"{left.strip()} -> no {dep} in frame"
    if "->" in s and "не пересекается с основным" in s:
        left, right = s.split("->", 1)
        dep = right.replace("не пересекается с основным", "").strip()
        return f"{left.strip()} -> no overlap with main ({dep})"
    if "->" in s and " и " in s and "не пересекаются" in s:
        left, right = s.split("->", 1)
        tail = right.replace("не пересекаются", "").strip()
        parts = [x.strip() for x in tail.split(" и ", 1)]
        if len(parts) == 2:
            return f"{left.strip()} -> no overlap: {parts[0]} & {parts[1]}"
    # OpenCV text rendering has poor unicode support; keep overlay strictly ASCII.
    out = re.sub(r"[^ -~]", " ", s)
    out = out.replace("?", " ")
    out = re.sub(r"\s+", " ", out).strip()
    return out or "warning condition failed"


def _draw_danger_crop(frame_bgr: np.ndarray, main_mask: np.ndarray, reasons: list[str]) -> np.ndarray:
    bb = _bbox_from_mask(main_mask)
    if bb is None:
        crop = frame_bgr.copy()
    else:
        x, y, bw, bh = bb
        padx = max(8, int(float(bw) * float(AN_CROP_PAD_RATIO)))
        pady = max(8, int(float(bh) * float(AN_CROP_PAD_RATIO)))
        x0 = max(0, x - padx)
        y0 = max(0, y - pady)
        x2 = min(frame_bgr.shape[1], x + bw + padx)
        y2 = min(frame_bgr.shape[0], y + bh + pady)
        crop = frame_bgr[y0:y2, x0:x2].copy()
        cv2.rectangle(
            crop,
            (max(0, x - x0), max(0, y - y0)),
            (max(0, x + bw - 1 - x0), max(0, y + bh - 1 - y0)),
            (30, 220, 255),
            2,
            lineType=cv2.LINE_AA,
        )
    bar_h = min(56, max(34, int(crop.shape[0] * 0.12)))
    overlay = crop.copy()
    cv2.rectangle(overlay, (0, 0), (crop.shape[1] - 1, bar_h), (0, 0, 180), thickness=-1)
    cv2.addWeighted(overlay, 0.65, crop, 0.35, 0.0, dst=crop)
    cv2.putText(
        crop,
        "DANGER",
        (10, int(bar_h * 0.7)),
        cv2.FONT_HERSHEY_SIMPLEX,
        0.8,
        (255, 255, 255),
        2,
        lineType=cv2.LINE_AA,
    )
    if reasons:
        txt = "; ".join(_reason_for_overlay(x) for x in reasons if str(x or "").strip())
        max_chars = max(20, int(crop.shape[1] / 9))
        if len(txt) > max_chars:
            txt = txt[: max_chars - 3] + "..."
        cv2.putText(
            crop,
            txt,
            (10, min(crop.shape[0] - 10, bar_h + 24)),
            cv2.FONT_HERSHEY_SIMPLEX,
            0.55,
            (245, 245, 245),
            1,
            lineType=cv2.LINE_AA,
        )
    return crop


@dataclass
class WarningItem:
    frame: int
    main_id: int
    reasons: list[str]
    image_url: str


def _find_video(folder: Path) -> Path | None:
    if not folder.is_dir():
        return None
    for p in folder.iterdir():
        if p.is_file() and p.suffix.lower() in VIDEO_EXTS and p.stem == "video":
            return p
    return None


def _safe_video_fps(video_path: Path) -> float:
    cap = cv2.VideoCapture(str(video_path))
    try:
        if not cap.isOpened():
            return 25.0
        fps = float(cap.get(cv2.CAP_PROP_FPS) or 0.0)
        if fps <= 0.1 or fps > 240.0:
            return 25.0
        return fps
    finally:
        cap.release()


def _build_warning_video(
    folder: str,
    folder_path: Path,
    video_path: Path,
    warnings: list[dict],
    main_id: int | None = None,
) -> dict[str, object]:
    analysis_dir = folder_path / "analysis"
    warn_dir = analysis_dir / "warnings"
    if main_id is None:
        out_name = "warnings_preview.mp4"
    else:
        out_name = f"warnings_preview_human_{int(main_id)}.mp4"
    out_path = analysis_dir / out_name

    items: list[tuple[int, Path]] = []
    for w in warnings:
        if main_id is not None:
            try:
                wid = int(w.get("main_id"))
            except Exception:
                continue
            if wid != int(main_id):
                continue
        try:
            fidx = int(w.get("frame", -1))
        except Exception:
            fidx = -1
        if fidx < 0:
            continue
        p: Path | None = None
        image_url = str(w.get("image_url", "") or "")
        if image_url:
            name = Path(urlparse(image_url).path).name
            if name:
                p0 = warn_dir / name
                if p0.is_file():
                    p = p0
        if p is None:
            p = warn_dir / f"warn_{fidx:06d}.jpg"
        if p.is_file():
            items.append((fidx, p))
    items.sort(key=lambda x: x[0])
    if not items:
        raise RuntimeError("No warning frames found. Run analysis first.")

    fps = _safe_video_fps(video_path)
    first = cv2.imread(str(items[0][1]), cv2.IMREAD_COLOR)
    if first is None or first.size == 0:
        raise RuntimeError("Cannot read first warning frame image.")
    h, w = int(first.shape[0]), int(first.shape[1])
    # libx264 + yuv420p requires even frame dimensions.
    w_enc = max(2, w - (w % 2))
    h_enc = max(2, h - (h % 2))
    tmp_dir = analysis_dir / "warnings_video_frames_tmp"
    if tmp_dir.exists():
        shutil.rmtree(tmp_dir, ignore_errors=True)
    tmp_dir.mkdir(parents=True, exist_ok=True)
    written = 0
    for _, p in items:
        frame = cv2.imread(str(p), cv2.IMREAD_COLOR)
        if frame is None or frame.size == 0:
            continue
        if frame.shape[0] != h_enc or frame.shape[1] != w_enc:
            frame = cv2.resize(frame, (w_enc, h_enc), interpolation=cv2.INTER_LINEAR)
        out_jpg = tmp_dir / f"frame_{written:06d}.jpg"
        ok = cv2.imwrite(str(out_jpg), frame, [int(cv2.IMWRITE_JPEG_QUALITY), 94])
        if ok:
            written += 1
    if written <= 0:
        raise RuntimeError("No valid warning frames to encode.")

    cmd = [
        "ffmpeg",
        "-y",
        "-framerate",
        f"{fps:.6f}",
        "-i",
        str(tmp_dir / "frame_%06d.jpg"),
        "-r",
        f"{fps:.6f}",
        "-pix_fmt",
        "yuv420p",
        "-c:v",
        "libx264",
        "-movflags",
        "+faststart",
        str(out_path),
    ]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True)
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)
    if proc.returncode != 0 or (not out_path.is_file()):
        err = (proc.stderr or proc.stdout or "").strip()
        raise RuntimeError(f"ffmpeg failed: {err[-600:]}")

    return {
        "video_url": f"/storage/folders/{folder}/analysis/{out_name}?ts={int(time.time())}",
        "frames_used": int(written),
        "fps": float(fps),
        "main_id": int(main_id) if main_id is not None else None,
    }


def _folder_row(folder: Path) -> dict[str, object]:
    json_path = folder / "data.json"
    video_path = _find_video(folder)
    mtime = folder.stat().st_mtime
    if json_path.is_file():
        mtime = max(mtime, json_path.stat().st_mtime)
    if video_path and video_path.is_file():
        mtime = max(mtime, video_path.stat().st_mtime)

    return {
        "name": folder.name,
        "has_video": bool(video_path and video_path.is_file()),
        "video_name": video_path.name if video_path else "",
        "has_json": json_path.is_file(),
        "updated_at": dt.datetime.fromtimestamp(mtime).isoformat(timespec="seconds"),
    }


def _list_folders() -> list[dict[str, object]]:
    _ensure_dirs()
    rows: list[dict[str, object]] = []
    for p in FOLDERS_ROOT.iterdir():
        if p.is_dir():
            rows.append(_folder_row(p))
    rows.sort(key=lambda x: str(x.get("updated_at", "")), reverse=True)
    return rows


def _processing_stats_for_folder(folder_name: str) -> dict[str, object]:
    _folder_path, payload, video_path = _load_folder_payload(folder_name)
    frames = payload.get("frames")
    frames_list = frames if isinstance(frames, list) else []
    frames_total = int(len(frames_list))
    frames_with_instances = 0
    instances_total = 0
    labels: set[str] = set()
    for fr in frames_list:
        if not isinstance(fr, dict):
            continue
        inst = fr.get("instances")
        if not isinstance(inst, list):
            continue
        valid = [x for x in inst if isinstance(x, dict)]
        if valid:
            frames_with_instances += 1
        instances_total += int(len(valid))
        for x in valid:
            lbl = str(x.get("prompt_label", "") or "").strip()
            if lbl:
                labels.add(lbl)

    proc = payload.get("web_samv_processing")
    proc_map = proc if isinstance(proc, dict) else {}
    scale_div = float(proc_map.get("scale_div", payload.get("fast_scale_div", 1.0)) or 1.0)
    fps_div = int(proc_map.get("fps_div", payload.get("fps_div", 1)) or 1)
    video_part_sec = float(proc_map.get("video_part_sec", 0.0) or 0.0)
    chunks_total = int(proc_map.get("chunks_total", 1) or 1)
    source_fps = float(proc_map.get("source_fps", payload.get("fps", 0.0)) or 0.0)
    processed_fps = float(proc_map.get("processed_fps", payload.get("fps", 0.0)) or 0.0)
    source_w = int(proc_map.get("source_width", payload.get("width", 0)) or 0)
    source_h = int(proc_map.get("source_height", payload.get("height", 0)) or 0)
    processed_w = int(proc_map.get("processed_width", payload.get("width", 0)) or 0)
    processed_h = int(proc_map.get("processed_height", payload.get("height", 0)) or 0)

    return {
        "folder": folder_name,
        "video_name": video_path.name if video_path and video_path.is_file() else "",
        "prompt": str(payload.get("prompt", "") or ""),
        "scale_div": float(scale_div),
        "fps_div": int(fps_div),
        "video_part_sec": float(video_part_sec),
        "chunks_total": int(chunks_total),
        "source_fps": float(source_fps),
        "processed_fps": float(processed_fps),
        "source_width": int(source_w),
        "source_height": int(source_h),
        "processed_width": int(processed_w),
        "processed_height": int(processed_h),
        "frames_written": int(payload.get("frames_written", 0) or 0),
        "frames_total": int(frames_total),
        "frames_with_instances": int(frames_with_instances),
        "instances_total": int(instances_total),
        "labels_count": int(len(labels)),
        "labels": sorted(labels),
        "elapsed_sec": float(payload.get("elapsed_sec", 0.0) or 0.0),
    }


def _processing_stats_all_folders() -> list[dict[str, object]]:
    rows: list[dict[str, object]] = []
    for item in _list_folders():
        name = str(item.get("name", "") or "")
        if not name:
            continue
        try:
            rows.append(_processing_stats_for_folder(name))
        except Exception:
            continue
    return rows


def _new_folder_name() -> str:
    return dt.datetime.now().strftime("run_%Y%m%d_%H%M%S")


def _effective_api_base() -> str:
    base = str(API_BASE or "").strip().rstrip("/")
    if base:
        return base
    # Auto mode without fixed IP: use configurable host hint (localhost by default).
    return f"http://{API_HOST_HINT}:{API_PORT_HINT}"


def _join_api(path: str) -> str:
    return urljoin(_effective_api_base() + "/", path.lstrip("/"))


def _http_json_get(path: str) -> dict:
    req = Request(_join_api(path), method="GET")
    with urlopen(req, timeout=120) as resp:
        raw = resp.read()
    out = json.loads(raw.decode("utf-8"))
    return out if isinstance(out, dict) else {}


def _http_bytes_get(path: str) -> bytes:
    req = Request(_join_api(path), method="GET")
    with urlopen(req, timeout=300) as resp:
        return resp.read()


def _api_health_status() -> dict[str, object]:
    base = _effective_api_base()
    try:
        j = _http_json_get("/health")
        ok = str(j.get("status", "")).lower() == "ok"
        return {"ok": ok, "api_base": base, "raw": j}
    except Exception as exc:
        return {"ok": False, "api_base": base, "error": str(exc)}


def _set_api_base(new_base: str) -> None:
    global API_BASE
    base = str(new_base or "").strip().rstrip("/")
    if not base:
        # Empty value resets to auto mode (WEB_SAMV_API_HOST + WEB_SAMV_API_PORT).
        API_BASE = ""
        return
    if not (base.startswith("http://") or base.startswith("https://")):
        raise ValueError("api_base must start with http:// or https://")
    API_BASE = base


def _multipart_body(fields: dict[str, str], file_field: str, file_name: str, file_bytes: bytes) -> tuple[bytes, str]:
    boundary = f"----samv-{uuid.uuid4().hex}"
    chunks: list[bytes] = []
    crlf = b"\r\n"
    for k, v in fields.items():
        chunks.append(f"--{boundary}".encode("utf-8"))
        chunks.append(f'Content-Disposition: form-data; name="{k}"'.encode("utf-8"))
        chunks.append(b"")
        chunks.append(str(v).encode("utf-8"))
    chunks.append(f"--{boundary}".encode("utf-8"))
    chunks.append(
        f'Content-Disposition: form-data; name="{file_field}"; filename="{file_name}"'.encode("utf-8")
    )
    chunks.append(b"Content-Type: application/octet-stream")
    chunks.append(b"")
    chunks.append(file_bytes)
    chunks.append(f"--{boundary}--".encode("utf-8"))
    chunks.append(b"")
    body = crlf.join(chunks)
    return body, boundary


def _video_meta_from_bytes(video_bytes: bytes, suffix: str = ".mp4") -> tuple[int, int, float] | None:
    tmp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(video_bytes)
            tmp_path = Path(tmp.name)
        cap = cv2.VideoCapture(str(tmp_path))
        try:
            if not cap.isOpened():
                return None
            w = int(cap.get(cv2.CAP_PROP_FRAME_WIDTH) or 0)
            h = int(cap.get(cv2.CAP_PROP_FRAME_HEIGHT) or 0)
            fps = float(cap.get(cv2.CAP_PROP_FPS) or 0.0)
            if w > 0 and h > 0:
                return w, h, fps
            return None
        finally:
            cap.release()
    except Exception:
        return None
    finally:
        try:
            if tmp_path is not None:
                tmp_path.unlink(missing_ok=True)
        except Exception:
            pass


def _video_duration_sec(video_bytes: bytes, suffix: str) -> float:
    ext = suffix if suffix.startswith(".") else f".{suffix}"
    tmp_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
            tmp.write(video_bytes)
            tmp_path = Path(tmp.name)
        cap = cv2.VideoCapture(str(tmp_path))
        if cap.isOpened():
            fps = float(cap.get(cv2.CAP_PROP_FPS) or 0)
            frames = float(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
            cap.release()
            if fps > 0 and frames > 0:
                return frames / fps
        proc = subprocess.run(
            [
                "ffprobe",
                "-v",
                "error",
                "-show_entries",
                "format=duration",
                "-of",
                "default=noprint_wrappers=1:nokey=1",
                str(tmp_path),
            ],
            capture_output=True,
            text=True,
        )
        if proc.returncode == 0:
            return float((proc.stdout or "0").strip() or 0)
    except Exception:
        return 0.0
    finally:
        try:
            if tmp_path is not None:
                tmp_path.unlink(missing_ok=True)
        except Exception:
            pass
    return 0.0


def _ffmpeg_cut_segment(video_bytes: bytes, suffix: str, start_sec: float, dur_sec: float) -> bytes | None:
    ext = suffix if suffix.startswith(".") else f".{suffix}"
    src_path: Path | None = None
    out_path: Path | None = None
    try:
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as src:
            src.write(video_bytes)
            src_path = Path(src.name)
        out_path = src_path.with_suffix(".part.mp4")
        cmd = [
            "ffmpeg",
            "-y",
            "-ss",
            f"{start_sec:.3f}",
            "-i",
            str(src_path),
            "-t",
            f"{dur_sec:.3f}",
            "-c:v",
            "libx264",
            "-pix_fmt",
            "yuv420p",
            "-an",
            str(out_path),
        ]
        proc = subprocess.run(cmd, capture_output=True, text=True)
        if proc.returncode != 0 or not out_path.is_file():
            return None
        return out_path.read_bytes()
    finally:
        for p in (src_path, out_path):
            if p is None:
                continue
            try:
                p.unlink(missing_ok=True)
            except Exception:
                pass


def _split_video_chunks(video_bytes: bytes, suffix: str, part_sec: float) -> list[bytes]:
    """part_sec <= 0 — одно видео целиком без нарезки."""
    sec = float(part_sec or 0)
    if sec <= 0:
        return [video_bytes]
    duration = _video_duration_sec(video_bytes, suffix)
    if duration <= 0 or duration <= sec + 0.05:
        return [video_bytes]
    out: list[bytes] = []
    start = 0.0
    while start < duration - 0.02:
        chunk = _ffmpeg_cut_segment(video_bytes, suffix, start, sec)
        if chunk:
            out.append(chunk)
        start += sec
    return out or [video_bytes]


def _video_part_sec_from_form(form) -> float:
    raw = str(form.getfirst("video_part_sec", "") or form.getfirst("part_sec", "") or "0").strip()
    try:
        sec = float(raw or 0)
    except Exception:
        sec = 0.0
    return max(0.0, sec)


def _merge_chunk_payloads(payloads: list[dict]) -> dict:
    if not payloads:
        return {}
    if len(payloads) == 1:
        return dict(payloads[0])
    merged: dict[str, object] = {
        "schema": payloads[0].get("schema", ""),
        "width": payloads[0].get("width", 0),
        "height": payloads[0].get("height", 0),
        "frames": [],
    }
    offset = 0
    for p in payloads:
        frames = p.get("frames")
        if not isinstance(frames, list):
            continue
        for fr in frames:
            if not isinstance(fr, dict):
                continue
            nf = dict(fr)
            try:
                nf["frame"] = int(nf.get("frame", 0)) + offset
            except Exception:
                nf["frame"] = offset
            merged["frames"].append(nf)
        offset = len(merged["frames"])
    return merged


def _transform_video_bytes_ffmpeg(
    video_bytes: bytes,
    *,
    suffix: str = ".mp4",
    scale_div: float = 1.0,
    fps_div: int = 1,
    src_fps: float = 0.0,
) -> bytes:
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as src_tmp:
        src_tmp.write(video_bytes)
        src_path = Path(src_tmp.name)
    with tempfile.NamedTemporaryFile(suffix=".mp4", delete=False) as dst_tmp:
        dst_path = Path(dst_tmp.name)
    vf_filters: list[str] = []
    sd = float(scale_div) if float(scale_div) > 0 else 1.0
    fd = int(fps_div) if int(fps_div) > 0 else 1
    if sd > 1.0001:
        # Keep output dimensions even for H.264.
        vf_filters.append(f"scale=trunc(iw/{sd}/2)*2:trunc(ih/{sd}/2)*2")
    if fd > 1:
        # Deterministic frame decimation: keep each N-th frame.
        vf_filters.append(f"select='not(mod(n,{fd}))'")
        vf_filters.append("setpts=N/FRAME_RATE/TB")
    cmd = [
        "ffmpeg",
        "-y",
        "-i",
        str(src_path),
        "-an",
        "-c:v",
        "libx264",
        "-preset",
        "veryfast",
        "-crf",
        "23",
        "-pix_fmt",
        "yuv420p",
        str(dst_path),
    ]
    if vf_filters:
        cmd[4:4] = ["-vf", ",".join(vf_filters)]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True)
        if proc.returncode != 0 or not dst_path.is_file():
            err = (proc.stderr or proc.stdout or "").strip()
            raise RuntimeError(f"ffmpeg transform failed: {err[-400:]}")
        return dst_path.read_bytes()
    finally:
        src_path.unlink(missing_ok=True)
        dst_path.unlink(missing_ok=True)


def _mask_bool_to_rle_row_major(mask_bool: np.ndarray) -> dict[str, object]:
    h, w = int(mask_bool.shape[0]), int(mask_bool.shape[1])
    flat = mask_bool.reshape(-1).astype(np.uint8, copy=False)
    n = int(flat.size)
    if n <= 0:
        return {"height": h, "width": w, "format": "rle_row_major", "order": "C", "counts": []}
    counts: list[int] = []
    cur = int(flat[0])
    run = 1
    for i in range(1, n):
        v = int(flat[i])
        if v == cur:
            run += 1
        else:
            counts.append(run)
            cur = v
            run = 1
    counts.append(run)
    if int(flat[0]) == 1:
        counts.insert(0, 0)
    return {"height": h, "width": w, "format": "rle_row_major", "order": "C", "counts": counts}


def _upscale_result_payload_inplace(payload: dict, target_w: int, target_h: int) -> None:
    src_w = int(payload.get("width", 0) or 0)
    src_h = int(payload.get("height", 0) or 0)
    if src_w <= 0 or src_h <= 0 or target_w <= 0 or target_h <= 0:
        return
    sx = float(target_w) / float(src_w)
    sy = float(target_h) / float(src_h)
    payload["width"] = int(target_w)
    payload["height"] = int(target_h)
    payload["fast_mode"] = "scaled"
    payload["fast_scale"] = float(sx if sx == sy else min(sx, sy))
    frames = payload.get("frames")
    if not isinstance(frames, list):
        return
    for fr in frames:
        if not isinstance(fr, dict):
            continue
        inst = fr.get("instances")
        if not isinstance(inst, list):
            continue
        for x in inst:
            if not isinstance(x, dict):
                continue
            bb = x.get("bbox_xywh")
            if isinstance(bb, list) and len(bb) >= 4:
                try:
                    bx, by, bw, bh = [float(bb[0]), float(bb[1]), float(bb[2]), float(bb[3])]
                    x["bbox_xywh"] = [
                        int(round(bx * sx)),
                        int(round(by * sy)),
                        int(round(bw * sx)),
                        int(round(bh * sy)),
                    ]
                except Exception:
                    pass
            cc = x.get("center_xy")
            if isinstance(cc, list) and len(cc) >= 2:
                try:
                    cx, cy = float(cc[0]), float(cc[1])
                    x["center_xy"] = [int(round(cx * sx)), int(round(cy * sy))]
                except Exception:
                    pass
            if "area_px" in x:
                try:
                    x["area_px"] = int(round(float(x.get("area_px", 0)) * sx * sy))
                except Exception:
                    pass
            enc = x.get("mask")
            if isinstance(enc, dict):
                mh = int(enc.get("height", 0) or 0)
                mw = int(enc.get("width", 0) or 0)
                m = _mask_from_rle_row_major(enc, mh, mw)
                if m is not None:
                    up = cv2.resize(m.astype(np.uint8), (target_w, target_h), interpolation=cv2.INTER_NEAREST).astype(bool)
                    x["mask"] = _mask_bool_to_rle_row_major(up)


def _start_video_job(video_name: str, video_bytes: bytes, prompt: str) -> str:
    fields = {
        "prompt": prompt,
        "alpha": "0.45",
        "visualize": "true",
        "enable_segmentation": "true",
        "draw_masks": "true",
        "draw_centers": "true",
        "export_stable_json": "false",
    }
    body, boundary = _multipart_body(fields, "file", video_name, video_bytes)
    req = Request(_join_api("/infer/video"), data=body, method="POST")
    req.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
    req.add_header("Content-Length", str(len(body)))
    with urlopen(req, timeout=300) as resp:
        raw = resp.read()
    out = json.loads(raw.decode("utf-8"))
    if not isinstance(out, dict):
        raise RuntimeError("Invalid /infer/video response")
    jid = str(out.get("job_id", "") or "")
    if not jid:
        raise RuntimeError("No job_id from /infer/video")
    return jid


def _wait_job_done(job_id: str) -> dict:
    started = time.time()
    while True:
        st = _http_json_get(f"/jobs/{job_id}")
        status = str(st.get("status", "") or "")
        if status == "done":
            return st
        if status == "failed":
            raise RuntimeError(str(st.get("error", "job failed")))
        if time.time() - started > JOB_TIMEOUT_SEC:
            raise TimeoutError(f"Job timeout after {JOB_TIMEOUT_SEC}s")
        time.sleep(max(0.2, POLL_INTERVAL_SEC))


def _save_processed_outputs(
    folder_name: str,
    result: dict,
    fallback_video: bytes,
    upscale_to_wh: tuple[int, int] | None = None,
    processing_meta: dict[str, object] | None = None,
) -> None:
    folder, json_path, _ = _folder_paths(folder_name)
    folder.mkdir(parents=True, exist_ok=True)

    preferred_urls = [
        str(result.get("download_data_stable_url") or "").strip(),
        str(result.get("download_data_url") or "").strip(),
    ]
    data_url = next((u for u in preferred_urls if u), "")
    raw_json: bytes | None = None
    if data_url:
        last_exc: Exception | None = None
        for u in preferred_urls:
            if not u:
                continue
            try:
                raw_json = _http_bytes_get(u)
                data_url = u
                last_exc = None
                break
            except Exception as exc:
                last_exc = exc
                continue
        if raw_json is None and last_exc is not None:
            raise last_exc
    if raw_json is not None:
        try:
            parsed = json.loads(raw_json.decode("utf-8"))
            if isinstance(parsed, dict) and upscale_to_wh is not None:
                _upscale_result_payload_inplace(parsed, int(upscale_to_wh[0]), int(upscale_to_wh[1]))
            if isinstance(parsed, dict) and isinstance(processing_meta, dict):
                parsed["web_samv_processing"] = dict(processing_meta)
                parsed["fast_scale_div"] = float(processing_meta.get("scale_div", 1.0) or 1.0)
                parsed["fps_div"] = int(processing_meta.get("fps_div", 1) or 1)
            json_path.write_text(json.dumps(parsed, ensure_ascii=False, indent=2), encoding="utf-8")
        except Exception:
            json_path.write_bytes(raw_json)
    else:
        json_path.write_text("{}", encoding="utf-8")

    video_url = str(result.get("download_video_url") or "")
    if video_url:
        vb = _http_bytes_get(video_url)
        ext = Path(urlparse(video_url).path).suffix.lower() or ".mp4"
        if ext not in VIDEO_EXTS:
            ext = ".mp4"
        (folder / f"video{ext}").write_bytes(vb)
    else:
        (folder / "video.mp4").write_bytes(fallback_video)


def _discovered_lan_ipv4() -> list[str]:
    out: list[str] = []
    seen: set[str] = set()

    def add(ip: str) -> None:
        if not ip or ip.startswith("127."):
            return
        if ip not in seen:
            seen.add(ip)
            out.append(ip)

    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as s:
            s.connect(("192.0.2.1", 1))
            add(s.getsockname()[0])
    except OSError:
        pass
    try:
        hn = socket.gethostname()
        for res in socket.getaddrinfo(hn, None, socket.AF_INET, socket.SOCK_STREAM):
            add(res[4][0])
    except OSError:
        pass
    return out


class _ReusableTCPServer(socketserver.TCPServer):
    allow_reuse_address = True


class Handler(http.server.SimpleHTTPRequestHandler):
    def end_headers(self) -> None:
        try:
            path_only = (self.path or "").split("?", 1)[0].lower()
            if path_only.endswith((".html", ".htm", ".js", ".css")):
                self.send_header("Cache-Control", "no-store, must-revalidate")
        except Exception:
            pass
        super().end_headers()

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        path = parsed.path.rstrip("/")
        if path == "/__web_samv_hint.json":
            _json_response(
                self,
                {
                    "ok": True,
                    "web_port": int(_RUNTIME.get("web_port", PORT)),
                    "api_port": int(_RUNTIME.get("api_port", API_PORT_HINT)),
                    "lan_ipv4": list(_RUNTIME.get("lan_ipv4", [])),
                },
            )
            return
        if path == "/api/folders":
            _json_response(self, {"ok": True, "root": str(FOLDERS_ROOT), "items": _list_folders()})
            return
        if path == "/api/status":
            st = _api_health_status()
            _json_response(self, {"ok": True, **st})
            return
        if path == "/api/analyzer/prompts":
            self._analyzer_prompts(parsed)
            return
        if path == "/api/analyzer/result":
            self._analyzer_result(parsed)
            return
        if path == "/api/folders/stats":
            self._folder_stats(parsed)
            return
        if path == "/api/folders/stats_all":
            self._folders_stats_all()
            return
        if path == "/api/inf/options":
            self._inf_options()
            return
        return super().do_GET()

    def do_POST(self) -> None:
        parsed = urlparse(self.path)
        path = parsed.path.rstrip("/")
        try:
            if path == "/api/folders/create":
                self._create_folder()
                return
            if path == "/api/folders/delete":
                self._delete_folder()
                return
            if path == "/api/folders/upload":
                self._upload_to_folder()
                return
            if path == "/api/folders/save_processed":
                self._save_processed_result()
                return
            if path == "/api/process_video":
                self._process_video()
                return
            if path == "/api/status/set":
                self._set_status_api_base()
                return
            if path == "/api/analyzer/run":
                self._analyzer_run()
                return
            if path == "/api/analyzer/video":
                self._analyzer_video()
                return
            if path == "/api/inf/add":
                self._inf_add()
                return
            if path == "/api/inf/delete":
                self._inf_delete()
                return
            if path == "/api/report/generate":
                self._report_generate()
                return
            _json_response(self, {"ok": False, "error": "Not found"}, code=404)
        except (BrokenPipeError, ConnectionAbortedError, ConnectionResetError):
            return

    def _read_json_body(self) -> dict:
        n = int(self.headers.get("Content-Length", "0") or 0)
        if n <= 0:
            return {}
        raw = self.rfile.read(n)
        try:
            parsed = json.loads(raw.decode("utf-8"))
            return parsed if isinstance(parsed, dict) else {}
        except Exception:
            return {}

    def _create_folder(self) -> None:
        _ensure_dirs()
        data = self._read_json_body()
        name = _safe_name(str(data.get("name", "") or ""))
        if not name:
            _json_response(self, {"ok": False, "error": "Invalid folder name"}, code=400)
            return
        folder, _, _ = _folder_paths(name)
        folder.mkdir(parents=True, exist_ok=True)
        _json_response(self, {"ok": True, "folder": name, "items": _list_folders()})

    def _delete_folder(self) -> None:
        _ensure_dirs()
        data = self._read_json_body()
        name = _safe_name(str(data.get("name", "") or ""))
        if not name:
            _json_response(self, {"ok": False, "error": "Invalid folder name"}, code=400)
            return
        folder, _, _ = _folder_paths(name)
        if not folder.is_dir():
            _json_response(self, {"ok": False, "error": "Folder not found"}, code=404)
            return
        shutil.rmtree(folder)
        _json_response(self, {"ok": True, "items": _list_folders()})

    def _upload_to_folder(self) -> None:
        _ensure_dirs()
        ctype = self.headers.get("Content-Type", "")
        if "multipart/form-data" not in ctype:
            _json_response(self, {"ok": False, "error": "Expected multipart/form-data"}, code=400)
            return

        try:
            form = cgi.FieldStorage(
                fp=self.rfile,
                headers=self.headers,
                environ={
                    "REQUEST_METHOD": "POST",
                    "CONTENT_TYPE": self.headers.get("Content-Type", ""),
                },
            )
        except Exception as exc:
            _json_response(self, {"ok": False, "error": f"Cannot parse multipart: {exc}"}, code=400)
            return

        folder_name = _safe_name(str(form.getfirst("folder", "") or ""))
        if not folder_name:
            _json_response(self, {"ok": False, "error": "Need folder field"}, code=400)
            return
        folder, json_path, _ = _folder_paths(folder_name)
        folder.mkdir(parents=True, exist_ok=True)

        if "video" in form and getattr(form["video"], "filename", ""):
            video_file = form["video"]
            video_name = str(video_file.filename or "")
            ext = Path(video_name).suffix.lower()
            if ext not in VIDEO_EXTS:
                _json_response(self, {"ok": False, "error": "Unsupported video extension"}, code=400)
                return
            for p in folder.iterdir():
                if p.is_file() and p.stem == "video" and p.suffix.lower() in VIDEO_EXTS:
                    p.unlink(missing_ok=True)
            video_dst = folder / f"video{ext}"
            with video_dst.open("wb") as f:
                f.write(video_file.file.read())

        if "json" in form and getattr(form["json"], "filename", ""):
            json_file = form["json"]
            json_name = str(json_file.filename or "")
            if Path(json_name).suffix.lower() != ".json":
                _json_response(self, {"ok": False, "error": "JSON file must have .json extension"}, code=400)
                return
            raw = json_file.file.read()
            try:
                parsed = json.loads(raw.decode("utf-8"))
            except Exception as exc:
                _json_response(self, {"ok": False, "error": f"Invalid JSON: {exc}"}, code=400)
                return
            json_path.write_text(json.dumps(parsed, ensure_ascii=False, indent=2), encoding="utf-8")

        _json_response(self, {"ok": True, "folder": folder_name, "items": _list_folders()})

    def _save_processed_result(self) -> None:
        """
        Автосоздание папки при обработке:
        POST /api/folders/save_processed
        {
          "folder": "site_001",
          "data": {...},                 # optional -> data.json
          "video_base64": "...",         # optional
          "video_ext": ".mp4"            # optional, default .mp4
        }
        """
        _ensure_dirs()
        data = self._read_json_body()
        folder_name = _safe_name(str(data.get("folder", "") or ""))
        if not folder_name:
            _json_response(self, {"ok": False, "error": "Need folder"}, code=400)
            return

        folder, json_path, _ = _folder_paths(folder_name)
        folder.mkdir(parents=True, exist_ok=True)

        payload = data.get("data")
        if isinstance(payload, dict):
            json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")

        video_b64 = str(data.get("video_base64", "") or "")
        if video_b64:
            ext = str(data.get("video_ext", ".mp4") or ".mp4").lower().strip()
            if not ext.startswith("."):
                ext = f".{ext}"
            if ext not in VIDEO_EXTS:
                _json_response(self, {"ok": False, "error": f"Unsupported video_ext: {ext}"}, code=400)
                return
            if "," in video_b64:
                video_b64 = video_b64.split(",", 1)[1]
            try:
                raw = base64.b64decode(video_b64, validate=False)
            except Exception as exc:
                _json_response(self, {"ok": False, "error": f"Invalid video_base64: {exc}"}, code=400)
                return
            for p in folder.iterdir():
                if p.is_file() and p.stem == "video" and p.suffix.lower() in VIDEO_EXTS:
                    p.unlink(missing_ok=True)
            (folder / f"video{ext}").write_bytes(raw)

        _json_response(self, {"ok": True, "folder": folder_name, "items": _list_folders()})

    def _process_video(self) -> None:
        _ensure_dirs()
        ctype = self.headers.get("Content-Type", "")
        if "multipart/form-data" not in ctype:
            _json_response(self, {"ok": False, "error": "Expected multipart/form-data"}, code=400)
            return

        try:
            form = cgi.FieldStorage(
                fp=self.rfile,
                headers=self.headers,
                environ={
                    "REQUEST_METHOD": "POST",
                    "CONTENT_TYPE": self.headers.get("Content-Type", ""),
                },
            )
        except Exception as exc:
            _json_response(self, {"ok": False, "error": f"Cannot parse multipart: {exc}"}, code=400)
            return

        if "video" not in form or not getattr(form["video"], "filename", ""):
            _json_response(self, {"ok": False, "error": "Need video file"}, code=400)
            return

        prompt = str(form.getfirst("prompt", "") or "").strip()
        # Backward compatibility: old checkbox flags.
        fast_x2 = str(form.getfirst("fast_x2", "0") or "0").strip().lower() in ("1", "true", "yes", "on")
        fps_half = str(form.getfirst("fps_half", "0") or "0").strip().lower() in ("1", "true", "yes", "on")
        raw_scale_div = str(form.getfirst("scale_div", "") or "").strip()
        raw_fps_div = str(form.getfirst("fps_div", "") or "").strip()
        try:
            scale_div = float(raw_scale_div) if raw_scale_div else (2.0 if fast_x2 else 1.0)
        except Exception:
            scale_div = 1.0
        try:
            fps_div = int(raw_fps_div) if raw_fps_div else (2 if fps_half else 1)
        except Exception:
            fps_div = 1
        if scale_div < 1.0:
            scale_div = 1.0
        if fps_div < 1:
            fps_div = 1
        video_part_sec = _video_part_sec_from_form(form)
        if not prompt:
            _json_response(self, {"ok": False, "error": "Need prompt"}, code=400)
            return

        vf = form["video"]
        video_name = str(vf.filename or "video.mp4")
        ext = Path(video_name).suffix.lower()
        if ext not in VIDEO_EXTS:
            _json_response(self, {"ok": False, "error": "Unsupported video extension"}, code=400)
            return
        video_bytes = vf.file.read()
        if not video_bytes:
            _json_response(self, {"ok": False, "error": "Video is empty"}, code=400)
            return

        meta = _video_meta_from_bytes(video_bytes, suffix=ext)
        original_wh = (int(meta[0]), int(meta[1])) if meta is not None else None
        original_fps = float(meta[2]) if meta is not None else 0.0

        folder_name = _new_folder_name()
        for _ in range(20):
            folder, _, _ = _folder_paths(folder_name)
            if not folder.exists():
                break
            folder_name = f"{_new_folder_name()}_{uuid.uuid4().hex[:4]}"

        try:
            chunks = _split_video_chunks(video_bytes, ext or ".mp4", video_part_sec)
            payloads: list[dict] = []
            job_ids: list[str] = []
            last_send_meta: tuple[int, int, float] | None = None

            for ci, chunk_bytes in enumerate(chunks):
                send_video_name = video_name
                send_video_bytes = chunk_bytes
                if len(chunks) > 1:
                    send_video_name = f"{Path(video_name).stem}_part{ci + 1:03d}{ext or '.mp4'}"
                if scale_div > 1.0 or fps_div > 1:
                    send_video_name = f"{Path(send_video_name).stem}.fast.mp4"
                    send_video_bytes = _transform_video_bytes_ffmpeg(
                        chunk_bytes,
                        suffix=ext or ".mp4",
                        scale_div=float(scale_div),
                        fps_div=int(fps_div),
                        src_fps=original_fps,
                    )
                last_send_meta = _video_meta_from_bytes(send_video_bytes, suffix=".mp4")

                job_id = _start_video_job(send_video_name, send_video_bytes, prompt)
                job_ids.append(job_id)
                _wait_job_done(job_id)
                result = _http_json_get(f"/jobs/{job_id}/result")

                if len(chunks) == 1:
                    _save_processed_outputs(
                        folder_name,
                        result,
                        fallback_video=video_bytes,
                        upscale_to_wh=original_wh if scale_div > 1.0 else None,
                        processing_meta={
                            "scale_div": float(scale_div),
                            "fps_div": int(fps_div),
                            "video_part_sec": float(video_part_sec),
                            "chunks_total": 1,
                            "source_fps": float(original_fps),
                            "processed_fps": float(last_send_meta[2]) if last_send_meta is not None else float(original_fps),
                            "source_width": int(original_wh[0]) if original_wh is not None else 0,
                            "source_height": int(original_wh[1]) if original_wh is not None else 0,
                            "processed_width": int(last_send_meta[0]) if last_send_meta is not None else int(original_wh[0]) if original_wh is not None else 0,
                            "processed_height": int(last_send_meta[1]) if last_send_meta is not None else int(original_wh[1]) if original_wh is not None else 0,
                        },
                    )
                    break

                data_url = str(result.get("download_data_url") or result.get("download_data_stable_url") or "")
                if data_url:
                    raw = _http_bytes_get(data_url)
                    parsed = json.loads(raw.decode("utf-8"))
                    if isinstance(parsed, dict):
                        if original_wh and scale_div > 1.0:
                            _upscale_result_payload_inplace(parsed, int(original_wh[0]), int(original_wh[1]))
                        payloads.append(parsed)

            if len(chunks) > 1:
                merged = _merge_chunk_payloads(payloads)
                merged["prompt"] = prompt
                if original_wh:
                    merged.setdefault("width", int(original_wh[0]))
                    merged.setdefault("height", int(original_wh[1]))
                merged["web_samv_processing"] = {
                    "scale_div": float(scale_div),
                    "fps_div": int(fps_div),
                    "video_part_sec": float(video_part_sec),
                    "chunks_total": len(chunks),
                    "source_fps": float(original_fps),
                    "processed_fps": float(last_send_meta[2]) if last_send_meta is not None else float(original_fps),
                    "source_width": int(original_wh[0]) if original_wh is not None else 0,
                    "source_height": int(original_wh[1]) if original_wh is not None else 0,
                    "processed_width": int(last_send_meta[0]) if last_send_meta is not None else int(original_wh[0]) if original_wh is not None else 0,
                    "processed_height": int(last_send_meta[1]) if last_send_meta is not None else int(original_wh[1]) if original_wh is not None else 0,
                }
                merged["fast_scale_div"] = float(scale_div)
                merged["fps_div"] = int(fps_div)
                folder, json_path, _ = _folder_paths(folder_name)
                folder.mkdir(parents=True, exist_ok=True)
                json_path.write_text(json.dumps(merged, ensure_ascii=False, indent=2), encoding="utf-8")
                (folder / "video.mp4").write_bytes(video_bytes)
        except Exception as exc:
            _json_response(self, {"ok": False, "error": f"Process failed: {exc}"}, code=500)
            return

        _json_response(
            self,
            {
                "ok": True,
                "folder": folder_name,
                "job_id": job_ids[-1] if job_ids else "",
                "job_ids": job_ids,
                "chunks_total": len(chunks),
                "video_part_sec": float(video_part_sec),
                "fast_x2": bool(scale_div == 2.0),
                "fps_half": bool(fps_div == 2),
                "scale_div": float(scale_div),
                "fps_div": int(fps_div),
                "items": _list_folders(),
            },
        )

    def _set_status_api_base(self) -> None:
        data = self._read_json_body()
        api_base = str(data.get("api_base", "") or "").strip()
        try:
            _set_api_base(api_base)
        except Exception as exc:
            _json_response(self, {"ok": False, "error": f"Invalid api_base: {exc}"}, code=400)
            return
        _json_response(self, {"ok": True, "api_base": _effective_api_base()})

    def _inf_options(self) -> None:
        opts = _inf_options()
        _json_response(self, {"ok": True, **opts, "root": str(INF_ROOT)})

    def _inf_add(self) -> None:
        data = self._read_json_body()
        kind = _norm_inf_kind(str(data.get("kind", "") or ""))
        value = _norm_inf_item(str(data.get("value", "") or ""))
        if not value:
            _json_response(self, {"ok": False, "error": "Need value"}, code=400)
            return
        try:
            path = _kind_to_inf_path(kind)
            defaults = _defaults_for_inf_path(path)
            current = _read_inf_list(path, defaults)
            if value not in current:
                current.append(value)
                _write_inf_list(path, current)
            _json_response(self, {"ok": True, **_inf_options()})
        except Exception as exc:
            _json_response(self, {"ok": False, "error": str(exc)}, code=400)

    def _inf_delete(self) -> None:
        data = self._read_json_body()
        kind = _norm_inf_kind(str(data.get("kind", "") or ""))
        value = _norm_inf_item(str(data.get("value", "") or ""))
        if not value:
            _json_response(self, {"ok": False, "error": "Need value"}, code=400)
            return
        try:
            path = _kind_to_inf_path(kind)
            defaults = _defaults_for_inf_path(path)
            current = _read_inf_list(path, defaults)
            current = [x for x in current if x != value]
            _write_inf_list(path, current)
            _json_response(self, {"ok": True, **_inf_options()})
        except Exception as exc:
            _json_response(self, {"ok": False, "error": str(exc)}, code=400)

    def _analyzer_prompts(self, parsed) -> None:
        qs = parse_qs(parsed.query or "")
        folder = _safe_name((qs.get("folder") or [""])[0])
        if not folder:
            _json_response(self, {"ok": False, "error": "Need folder"}, code=400)
            return
        try:
            _, payload, _ = _load_folder_payload(folder)
            prompts = _all_prompts(payload)
            _json_response(self, {"ok": True, "folder": folder, "prompts": prompts})
        except Exception as exc:
            _json_response(self, {"ok": False, "error": str(exc)}, code=400)

    def _analyzer_result(self, parsed) -> None:
        qs = parse_qs(parsed.query or "")
        folder = _safe_name((qs.get("folder") or [""])[0])
        if not folder:
            _json_response(self, {"ok": False, "error": "Need folder"}, code=400)
            return
        try:
            folder_path, _, _ = _load_folder_payload(folder)
        except Exception as exc:
            _json_response(self, {"ok": False, "error": str(exc)}, code=400)
            return
        result_path = folder_path / "analysis" / "analyzer_result.json"
        if not result_path.is_file():
            _json_response(self, {"ok": True, "exists": False, "folder": folder})
            return
        try:
            data = json.loads(result_path.read_text(encoding="utf-8"))
            if not isinstance(data, dict):
                raise ValueError("Invalid analyzer_result.json")
        except Exception as exc:
            _json_response(self, {"ok": False, "error": f"Cannot read analyzer_result.json: {exc}"}, code=500)
            return
        preview_path = folder_path / "analysis" / "warnings_preview.mp4"
        preview_url = ""
        if preview_path.is_file():
            preview_url = f"/storage/folders/{folder}/analysis/warnings_preview.mp4?ts={int(time.time())}"
        _json_response(self, {"ok": True, "exists": True, "folder": folder, "preview_video_url": preview_url, **data})

    def _folder_stats(self, parsed) -> None:
        qs = parse_qs(parsed.query or "")
        folder = _safe_name((qs.get("folder") or [""])[0])
        if not folder:
            _json_response(self, {"ok": False, "error": "Need folder"}, code=400)
            return
        try:
            stats = _processing_stats_for_folder(folder)
        except Exception as exc:
            _json_response(self, {"ok": False, "error": str(exc)}, code=400)
            return
        _json_response(self, {"ok": True, **stats})

    def _folders_stats_all(self) -> None:
        try:
            rows = _processing_stats_all_folders()
        except Exception as exc:
            _json_response(self, {"ok": False, "error": str(exc)}, code=500)
            return
        _json_response(self, {"ok": True, "items": rows})

    def _analyzer_run(self) -> None:
        data = self._read_json_body()
        folder = _safe_name(str(data.get("folder", "") or ""))
        main_prompt = str(data.get("main_prompt", "") or "").strip()
        deps_raw = data.get("linked_prompts", [])
        if not folder or not main_prompt:
            _json_response(self, {"ok": False, "error": "Need folder and main_prompt"}, code=400)
            return
        linked_prompts: list[str] = []
        if isinstance(deps_raw, list):
            for x in deps_raw:
                s = str(x or "").strip()
                if s and s != main_prompt and s not in linked_prompts:
                    linked_prompts.append(s)
        linked_prompts = linked_prompts[:2]
        try:
            folder_path, payload, video_path = _load_folder_payload(folder)
        except Exception as exc:
            _json_response(self, {"ok": False, "error": str(exc)}, code=400)
            return
        if video_path is None or not video_path.is_file():
            _json_response(self, {"ok": False, "error": "video.* not found in folder"}, code=400)
            return

        h = int(payload.get("height", 0) or 0)
        w = int(payload.get("width", 0) or 0)
        if h <= 0 or w <= 0:
            _json_response(self, {"ok": False, "error": "Invalid width/height in data.json"}, code=400)
            return
        frames = payload.get("frames")
        if not isinstance(frames, list):
            _json_response(self, {"ok": False, "error": "Invalid frames in data.json"}, code=400)
            return

        analysis_dir = folder_path / "analysis"
        warn_dir = analysis_dir / "warnings"
        warn_dir.mkdir(parents=True, exist_ok=True)
        for old in warn_dir.glob("warn_*.jpg"):
            old.unlink(missing_ok=True)

        cap = cv2.VideoCapture(str(video_path))
        if not cap.isOpened():
            _json_response(self, {"ok": False, "error": f"Cannot open video: {video_path.name}"}, code=500)
            return
        warnings_out: list[WarningItem] = []
        frames_checked = 0
        frames_main_present = 0
        try:
            for fr in frames:
                if not isinstance(fr, dict):
                    continue
                fidx = int(fr.get("frame", -1))
                if fidx < 0:
                    continue
                inst = fr.get("instances")
                if not isinstance(inst, list):
                    continue
                inst2 = [x for x in inst if isinstance(x, dict)]
                frames_checked += 1
                main_items = _instance_group_by_label(inst2, main_prompt, h, w)
                if not main_items:
                    continue
                frames_main_present += 1
                frame_warnings: list[tuple[int, np.ndarray, list[str]]] = []

                # Проверяем каждого main_id отдельно.
                for main_id, main_mask in main_items:
                    local_reasons: list[str] = []
                    dep_hits: list[list[tuple[int, np.ndarray]]] = []

                    for dep in linked_prompts:
                        dep_items = _instance_group_by_label(inst2, dep, h, w)
                        if not dep_items:
                            local_reasons.append(
                                f"{main_prompt}_id:{main_id} -> нет {dep} на кадре"
                            )
                            dep_hits.append([])
                            continue
                        inter = [
                            (dep_id, dep_mask)
                            for dep_id, dep_mask in dep_items
                            if _mask_intersects(main_mask, dep_mask)
                        ]
                        dep_hits.append(inter)
                        # Каждый связанный промт обязан входить в маску текущего main_id.
                        # Иначе это WARNING, даже если второй связанный промт присутствует.
                        if not inter:
                            local_reasons.append(
                                f"{main_prompt}_id:{main_id} -> {dep} не пересекается с основным"
                            )

                    if (
                        len(linked_prompts) == 2
                        and len(dep_hits) >= 2
                        and dep_hits[0]
                        and dep_hits[1]
                    ):
                        ok_pair = False
                        for _, d1 in dep_hits[0]:
                            for _, d2 in dep_hits[1]:
                                if _mask_overlap_ok(d1, d2):
                                    ok_pair = True
                                    break
                            if ok_pair:
                                break
                        if not ok_pair:
                            local_reasons.append(
                                f"{main_prompt}_id:{main_id} -> "
                                f"{linked_prompts[0]} и {linked_prompts[1]} не пересекаются"
                            )

                    if local_reasons:
                        frame_warnings.append((int(main_id), main_mask, local_reasons))

                if not frame_warnings:
                    continue
                cap.set(cv2.CAP_PROP_POS_FRAMES, float(fidx))
                ok, frame = cap.read()
                if not ok or frame is None:
                    continue
                if frame.shape[0] != h or frame.shape[1] != w:
                    frame = cv2.resize(frame, (w, h), interpolation=cv2.INTER_LINEAR)
                for main_id, main_mask, main_reasons in frame_warnings:
                    out_img = _draw_danger_crop(frame, main_mask, main_reasons)
                    out_name = f"warn_{fidx:06d}_human_{main_id}.jpg"
                    out_path = warn_dir / out_name
                    cv2.imwrite(str(out_path), out_img, [int(cv2.IMWRITE_JPEG_QUALITY), 92])
                    image_url = f"/storage/folders/{folder}/analysis/warnings/{out_name}"
                    warnings_out.append(
                        WarningItem(frame=fidx, main_id=int(main_id), reasons=main_reasons, image_url=image_url)
                    )
        finally:
            cap.release()

        summary = {
            "schema": "samv_mask_analyzer_v1",
            "folder": folder,
            "main_prompt": main_prompt,
            "linked_prompts": linked_prompts,
            "frames_checked": int(frames_checked),
            "frames_with_main": int(frames_main_present),
            "warnings_count": int(len(warnings_out)),
            "warnings": [
                {"frame": w.frame, "main_id": w.main_id, "reasons": w.reasons, "image_url": w.image_url}
                for w in warnings_out
            ],
            "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
        }
        (analysis_dir / "analyzer_result.json").write_text(
            json.dumps(summary, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        _json_response(self, {"ok": True, **summary})

    def _analyzer_video(self) -> None:
        data = self._read_json_body()
        folder = _safe_name(str(data.get("folder", "") or ""))
        if not folder:
            _json_response(self, {"ok": False, "error": "Need folder"}, code=400)
            return
        raw_main_id = data.get("main_id", None)
        main_id: int | None = None
        if raw_main_id is not None and str(raw_main_id).strip() != "":
            try:
                main_id = int(raw_main_id)
            except Exception:
                _json_response(self, {"ok": False, "error": "main_id must be integer"}, code=400)
                return
        try:
            folder_path, _, video_path = _load_folder_payload(folder)
        except Exception as exc:
            _json_response(self, {"ok": False, "error": str(exc)}, code=400)
            return
        if video_path is None or not video_path.is_file():
            _json_response(self, {"ok": False, "error": "video.* not found in folder"}, code=400)
            return
        result_path = folder_path / "analysis" / "analyzer_result.json"
        if not result_path.is_file():
            _json_response(self, {"ok": False, "error": "Run analysis first"}, code=400)
            return
        try:
            payload = json.loads(result_path.read_text(encoding="utf-8"))
            warnings = payload.get("warnings")
            if not isinstance(warnings, list):
                warnings = []
            built = _build_warning_video(folder, folder_path, video_path, warnings, main_id=main_id)
        except Exception as exc:
            _json_response(self, {"ok": False, "error": f"Build video failed: {exc}"}, code=500)
            return
        _json_response(self, {"ok": True, "folder": folder, **built})

    def _report_generate(self) -> None:
        data = self._read_json_body()
        folder = _safe_name(str(data.get("folder", "") or ""))
        if not folder:
            _json_response(self, {"ok": False, "error": "Need folder"}, code=400)
            return
        frame = int(data.get("frame", -1) or -1)
        image_url = str(data.get("image_url", "") or "")
        items = _norm_report_items(data.get("items"))
        if not items:
            # Backward compatibility: one global pair
            building = _norm_inf_item(str(data.get("building", "") or ""))
            contractor = _norm_inf_item(str(data.get("contractor", "") or ""))
            if building and contractor:
                contract = _norm_inf_item(str(data.get("contract", "") or ""))
                items = [
                    {
                        "human_id": "0",
                        "building": building,
                        "contractor": contractor,
                        "contract": contract,
                        "reasons": [],
                    }
                ]
        if not items:
            _json_response(self, {"ok": False, "error": "Need items with per-human selections"}, code=400)
            return
        formats = _norm_report_formats(data.get("formats"))
        try:
            out = _generate_report(folder, frame, image_url, items, formats=formats)
            _json_response(self, out)
        except Exception as exc:
            _json_response(self, {"ok": False, "error": f"Report failed: {exc}"}, code=500)


def _addr_in_use(exc: OSError) -> bool:
    if getattr(exc, "winerror", None) == 10048:
        return True
    return exc.errno == errno.EADDRINUSE


def _bind_server() -> tuple[_ReusableTCPServer, int]:
    last_exc: OSError | None = None
    for i in range(MAX_PORT_TRIES):
        port = PORT + i
        try:
            return _ReusableTCPServer(("0.0.0.0", port), Handler), port
        except OSError as exc:
            last_exc = exc
            if _addr_in_use(exc):
                continue
            raise
    raise OSError(f"Cannot bind port range {PORT}-{PORT + MAX_PORT_TRIES - 1}: {last_exc!r}")


def main() -> None:
    _ensure_dirs()
    os.chdir(str(ROOT))
    httpd, port = _bind_server()
    _RUNTIME["web_port"] = port
    _RUNTIME["lan_ipv4"] = _discovered_lan_ipv4()
    _RUNTIME["api_port"] = API_PORT_HINT
    print("--- WEB_samv ---")
    print(f"Root: {ROOT}")
    print(f"Data: {FOLDERS_ROOT}")
    print(f"Open: http://127.0.0.1:{port}/")
    lan = list(_RUNTIME.get("lan_ipv4", []))
    if lan:
        print("LAN:")
        for ip in lan:
            print(f"  http://{ip}:{port}/")
        print(f"API hint: http://{lan[0]}:{API_PORT_HINT}/")
    print("Stop: Ctrl+C")
    print()
    with httpd:
        try:
            httpd.serve_forever()
        except KeyboardInterrupt:
            print("\nStopped.")


if __name__ == "__main__":
    main()
